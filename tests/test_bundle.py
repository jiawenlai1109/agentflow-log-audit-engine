"""Bundle 与 ingest 的契约测试（M2-1）。

重点不是"能不能读文件"，而是三条纪律有没有被写死：
文档永不进 tables、原件不可变且带 sha256、坏文件在严格模式下必须整体失败。
"""

import json

import pandas as pd
import pytest

from agentflow.core.bundle import Bundle, BundleError
from agentflow.core.ingest import IngestError, build_bundle


def write(path, text: str):
    path.write_text(text, encoding="utf-8")
    return path


@pytest.fixture()
def sample_files(tmp_path):
    write(tmp_path / "sales.csv", "日期,销售额,主机\n2024-01-01,100,host-a\n2024-01-02,200,host-b\n")
    write(tmp_path / "assets.tsv", "主机\t域\t负责人\nhost-a\t生产\t张三\nhost-b\t测试\t李四\n")
    write(
        tmp_path / "alerts.json",
        json.dumps([{"主机": "host-a", "次数": 3}, {"主机": "host-b", "次数": 5}]),
    )
    write(
        tmp_path / "events.jsonl",
        '{"主机":"host-a","动作":"登录成"}\n{"主机":"host-b","动作":"登录失败"}\n',
    )
    write(tmp_path / "readme.md", "# 说明\n这份文件不是数据。\n" + "x" * 3000)
    write(tmp_path / "app.log", "2024-01-01 00:00:00 INFO started\nplain text line\n")
    return tmp_path


def test_tabular_formats_become_tables(sample_files, tmp_path):
    bundle = build_bundle(
        [sample_files / name for name in ("sales.csv", "assets.tsv", "alerts.json", "events.jsonl")],
        tmp_path / "bundle_strict",
    )
    assert bundle.table_ids() == ["t1", "t2", "t3", "t4"]
    assert [t.row_count for t in bundle.tables] == [2, 2, 2, 2]
    assert bundle.tables[0].columns == ["日期", "销售额", "主机"]
    # 归一化产物必须可直接被 pandas 读回
    frame = pd.read_csv(bundle.tables[1].path)
    assert list(frame.columns) == ["主机", "域", "负责人"]


def test_documents_never_enter_tables(sample_files, tmp_path):
    """I1 的地基：txt/log/md 只能作证据，不能出现在可聚合清单里。"""
    bundle = build_bundle(
        [sample_files / "sales.csv", sample_files / "readme.md", sample_files / "app.log"],
        tmp_path / "bundle_mixed",
    )
    assert bundle.table_ids() == ["t1"]
    assert {d.source_file for d in bundle.documents} == {"readme.md", "app.log"}
    assert all("x" * 3000 not in t.columns for t in bundle.tables)
    assert bundle.documents[0].preview.startswith("# 说明")
    # 文档正文长度不该把预览撑爆
    assert len(bundle.documents[0].preview) <= 2000


def test_documents_only_bundle_is_rejected(sample_files, tmp_path):
    with pytest.raises(BundleError):
        build_bundle([sample_files / "readme.md", sample_files / "app.log"], tmp_path / "bundle_docs_only")


def test_originals_are_copied_and_hashed(sample_files, tmp_path):
    bundle = build_bundle([sample_files / "sales.csv"], tmp_path / "bundle_copy")
    table = bundle.tables[0]
    source = bundle.root / "sources" / "sales.csv"
    assert source.exists() and Path_eq(table.source_path, source)
    assert len(table.sha256) == 16
    # 原件被改动不应影响已入包的副本
    sample_files.joinpath("sales.csv").write_text("日期,销售额\n2099-01-01,1\n", encoding="utf-8")
    assert "2024-01-01" in source.read_text(encoding="utf-8-sig")


def Path_eq(left, right) -> bool:
    from pathlib import Path

    return Path(left).resolve() == Path(right).resolve()


def test_manifest_round_trip(sample_files, tmp_path):
    root = tmp_path / "bundle_manifest"
    bundle = build_bundle([sample_files / "sales.csv", sample_files / "readme.md"], root)
    bundle.write()
    reloaded = Bundle.load(root)
    assert reloaded.table_ids() == bundle.table_ids()
    assert {d.source_file for d in reloaded.documents} == {"readme.md"}
    assert reloaded.tables[0].sha256 == bundle.tables[0].sha256
    assert reloaded.primary.id == "t1"


def test_join_candidates_find_the_shared_key(sample_files, tmp_path):
    bundle = build_bundle(
        [sample_files / "sales.csv", sample_files / "assets.tsv"], tmp_path / "bundle_join"
    )
    candidates = bundle.join_candidates()
    hit = next((item for item in candidates if item["column"] == "主机"), None)
    assert hit and hit["left"] == "t1" and hit["right"] == "t2"
    assert hit["usable"] is True and hit["overlap"] == 1.0


def test_unshared_columns_produce_no_candidate(sample_files, tmp_path):
    other = sample_files / "other.csv"
    write(other, "城市,人口\n上海,2400\n北京,2100\n")
    bundle = build_bundle([sample_files / "sales.csv", other], tmp_path / "bundle_nojoin")
    assert bundle.join_candidates() == []


def test_non_object_jsonl_degrades_to_document(sample_files, tmp_path):
    messy = sample_files / "messy.jsonl"
    write(messy, '{"a":1}\nnot a json line\n')
    bundle = build_bundle([sample_files / "sales.csv", messy], tmp_path / "bundle_messy")
    assert bundle.table_ids() == ["t1"]
    assert [d.source_file for d in bundle.documents] == ["messy.jsonl"]


def test_empty_but_well_formed_table_is_accepted(tmp_path):
    """0 行、列齐全是合法输入——空结果语义反转要靠它。"""
    zero = write(tmp_path / "zero.csv", "日期,销售额\n")
    bundle = build_bundle([zero], tmp_path / "bundle_zero")
    assert bundle.tables[0].row_count == 0
    assert bundle.tables[0].columns == ["日期", "销售额"]


def test_unknown_extension_rejected_strict_and_recorded_lenient(sample_files, tmp_path):
    mystery = write(sample_files / "mystery.bin", "\x00\x01binary")
    files = [sample_files / "sales.csv", mystery]
    with pytest.raises(ValueError, match="不支持的文件类型"):
        build_bundle(files, tmp_path / "bundle_strict_reject")
    lenient = build_bundle(files, tmp_path / "bundle_lenient", strict=False)
    assert lenient.table_ids() == ["t1"]
    assert lenient.skipped and lenient.skipped[0]["file"] == "mystery.bin"


def test_optional_dependency_failure_carries_install_hint(tmp_path):
    """缺 openpyxl 时报"装什么"，而不是抛一串 ImportError 让人猜。"""
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        pass
    else:
        pytest.skip("环境已装 openpyxl，本用例只在缺依赖时成立")
    fake = write(tmp_path / "book.xlsx", "whatever")
    with pytest.raises(IngestError) as error:
        build_bundle([fake], tmp_path / "bundle_xlsx")
    assert "openpyxl" in str(error.value) or "openpyxl" in error.value.hint
    assert "pip install" in error.value.hint


def test_json_columnar_shape_becomes_table(tmp_path):
    frame_file = write(tmp_path / "cols.json", json.dumps({"日期": ["2024-01-01"], "值": [1], "组": ["a"]}))
    bundle = build_bundle([frame_file], tmp_path / "bundle_cols")
    assert bundle.tables[0].columns == ["日期", "值", "组"]


def test_json_scalar_object_becomes_single_row(tmp_path):
    meta = write(tmp_path / "meta.json", json.dumps({"窗口": "5m", "阈值": 10}))
    bundle = build_bundle([meta], tmp_path / "bundle_meta")
    assert bundle.tables[0].row_count == 1
    assert set(bundle.tables[0].columns) == {"窗口", "阈值"}


def test_empty_array_is_rejected(tmp_path):
    empty = write(tmp_path / "empty.json", "[]")
    with pytest.raises(ValueError, match="JSON 数组为空"):
        build_bundle([empty], tmp_path / "bundle_empty")


def test_summary_states_the_discipline(sample_files, tmp_path):
    bundle = build_bundle([sample_files / "sales.csv", sample_files / "readme.md"], tmp_path / "bundle_sum")
    text = bundle.summary()
    assert "t1=sales.csv(2行×3列)" in text
    assert "不作数字来源" in text


def test_table_lookup_by_ref_or_name(sample_files, tmp_path):
    bundle = build_bundle([sample_files / "sales.csv"], tmp_path / "bundle_ref")
    assert bundle.table("t1").source_file == "sales.csv"
    assert bundle.table("sales.csv").id == "t1"
    with pytest.raises(BundleError):
        bundle.table("t9")


def test_missing_manifest_is_not_a_bundle(tmp_path):
    with pytest.raises(BundleError):
        Bundle.load(tmp_path)
